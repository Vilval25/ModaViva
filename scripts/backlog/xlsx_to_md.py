"""Genera docs/backlog/*.md a partir del Excel del backlog.

Se usa después de revisar el Excel al cerrar un hito. Sobrescribe los .md, así
que antes hay que haber llevado al Excel los estados de los criterios.

    python scripts/backlog/xlsx_to_md.py "<ruta>/Backlog ModaViva.xlsx" docs/backlog
"""
import re, sys, unicodedata
from pathlib import Path
import openpyxl

XLSX = Path(sys.argv[1])
OUT = Path(sys.argv[2])
OUT.mkdir(parents=True, exist_ok=True)

wb = openpyxl.load_workbook(XLSX)
ws = wb["PLANIFICACION"]


def slug(text):
    t = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "-", t).strip("-")


def bullets(text):
    """Convierte '• a\n• b' en lista Markdown; deja el resto como párrafo."""
    if text is None:
        return "—"
    text = str(text).strip()
    lines = [l.strip() for l in text.split("\n") if l.strip()]
    if lines and all(l.startswith("•") for l in lines):
        return "\n".join("- " + l.lstrip("• ").strip() for l in lines)
    return text


stories = []
sprint = epic = None
for r in range(3, ws.max_row + 1):
    b, c, d, h, i = (ws.cell(r, col).value for col in (2, 3, 4, 8, 9))
    if b:
        sprint = str(b)
    if c:
        epic = str(c)
    if d:
        head, _, story = str(d).partition("\n\n")
        sid, _, title = head.partition(" - ")
        sid = sid.strip()
        code = sid.split(" ")[0]
        cell = lambda col: ws.cell(r, col).value
        stories.append({
            "code": code,
            "kind": "Historia técnica" if "técnica" in sid else "Historia de usuario",
            "title": title.strip(),
            "story": story.strip(),
            "sprint": sprint,
            "epic": epic,
            "desc": cell(5), "pantallas": cell(6), "componentes": cell(7),
            "tecnologias": cell(11), "semanas": cell(12), "encargado": cell(13),
            "dependencias": cell(14), "requisitos": cell(15), "herramientas": cell(16),
            "criterios": [],
        })
    if h:
        stories[-1]["criterios"].append((str(h).strip(), (i or "Pendiente").strip()))

DRAFT = ("> ⚠️ **Borrador.** Las historias del Sprint 2 y 3 son una base y se revisarán "
         "al cerrar el Hito 1. No implementar sin confirmar que el texto es definitivo.\n\n")

index_rows = []
for s in stories:
    sprint_name, _, sprint_weeks = s["sprint"].partition("\n\n")
    epic_name = s["epic"].split("\n\n")[0]
    epic_desc = s["epic"].partition("\n\n")[2].strip()
    draft = sprint_name.strip() != "Sprint 1"
    fname = f"{s['code']}-{slug(s['title'])}.md"
    s["file"] = fname

    crit = []
    for n, (text, state) in enumerate(s["criterios"], 1):
        box = "x" if state == "Listo" else " "
        extra = " _(en progreso)_" if state == "En progreso" else ""
        crit.append(f"- [{box}] **CA-{n:02d}** {text}{extra}")

    md = f"""# {s['code']} — {s['title']}

{DRAFT if draft else ''}| Campo | Detalle |
| :--- | :--- |
| **Tipo** | {s['kind']} |
| **Épica** | {epic_name} |
| **Sprint** | {sprint_name.strip()} ({sprint_weeks.strip()}) |
| **Semanas** | {s['semanas'] or '—'} |
| **Encargado** | {s['encargado'] or '—'} |
| **Dependencias** | {s['dependencias'] or '—'} |
| **Requisitos del curso** | {s['requisitos'] or '—'} |

> {s['story']}

## Descripción funcional

{bullets(s['desc'])}

## Pantallas

{bullets(s['pantallas'])}

## Componentes UI

{bullets(s['componentes'])}

## Criterios de aceptación

{chr(10).join(crit)}

## Tecnologías

{bullets(s['tecnologias'])}

**Herramientas:** {s['herramientas'] or '—'}
"""
    (OUT / fname).write_text(md, encoding="utf-8")
    done = sum(1 for _, st in s["criterios"] if st == "Listo")
    index_rows.append((sprint_name.strip(), epic_name, s, done, draft))

# Índice
lines = []
current_sprint = None
for sprint_name, epic_name, s, done, draft in index_rows:
    if sprint_name != current_sprint:
        current_sprint = sprint_name
        weeks = s["sprint"].partition("\n\n")[2].strip()
        note = " — _borrador, se revisa al cerrar el Hito 1_" if draft else ""
        lines.append(f"\n### {sprint_name} ({weeks}){note}\n")
        lines.append("| Historia | Épica | Encargado | Semanas | Criterios |")
        lines.append("| :--- | :--- | :--- | :--- | :---: |")
    total = len(s["criterios"])
    lines.append(f"| [{s['code']} — {s['title']}]({s['file']}) | {epic_name.split('.')[0]} | "
                 f"{s['encargado']} | {s['semanas']} | {done}/{total} |")

epics = []
for _, epic_name, s, _, _ in index_rows:
    if epic_name not in [e[0] for e in epics]:
        epics.append((epic_name, s["epic"].partition("\n\n")[2].strip()))
epics.sort(key=lambda e: int(e[0][1:].split(".")[0]))

readme = f"""# Backlog de ModaViva

**Fuente de verdad del backlog durante el desarrollo.** El Excel del backlog
(entregable del curso, se guarda fuera del repositorio) se actualiza a partir de
estos archivos al cerrar cada hito; entre hitos no se edita.

## Cómo se usa

- Cada historia tiene su archivo. Sus criterios son una checklist: al implementar y
  **comprobar en un dispositivo** un criterio, se marca `[x]` en el mismo commit.
- Los criterios se citan por código, p. ej. `HU-01 CA-03`, en commits y PRs.
- Al marcar criterios, actualizar la columna **Criterios** de este índice (`hechos/total`).
- Una historia está terminada cuando cumple la [Definition of Done](SUPUESTOS.md#definition-of-done).
- Las historias del Sprint 2 y 3 son borrador hasta que se revise el Excel al cerrar el Hito 1.

## Hitos

- **Hito 1** (fin de SEM 7): cuenta + catálogo navegable con carrito.
- **Hito 2** (fin de SEM 11): probador virtual con IA.
- **Hito 3** (fin de SEM 15): compra, pago, postventa, analítica, asistencia con IA y publicación.

## Épicas

{chr(10).join(f"- **{n}** — {d}" for n, d in epics)}

## Historias
{chr(10).join(lines)}

Ver también: [Supuestos, DoR, DoD y requisitos del curso](SUPUESTOS.md).
"""
(OUT / "README.md").write_text(readme, encoding="utf-8")

# SUPUESTOS + REQUISITOS
sup = wb["SUPUESTOS"]
out = ["# Supuestos, definiciones y requisitos del curso\n"]
skip_section = False
for r in range(4, sup.max_row + 1):
    b, c = sup.cell(r, 2).value, sup.cell(r, 3).value
    if b and not c:
        skip_section = b.startswith("Cambios respecto")
        if skip_section:
            continue
        title = b
        if b.startswith("Definición de Listo"):
            title = "Definition of Ready\n\n_" + b + "_"
        elif b.startswith("Definición de Terminado"):
            title = "Definition of Done\n\n_" + b + "_"
        out.append(f"\n## {title}\n")
    elif b and c and not skip_section:
        if b in ("Definition of Ready", "Definition of Done"):
            out.append(bullets(c) + "\n")
        else:
            out.append(f"- **{b}:** {c}")
req = wb["REQUISITOS"]
out.append("\n## Requisitos del curso\n")
out.append("| N° | Requisito | Historias que lo cubren |")
out.append("| :---: | :--- | :--- |")
for r in range(6, req.max_row + 1):
    n, name, hus = (req.cell(r, col).value for col in (2, 3, 4))
    if name:
        out.append(f"| {int(n)} | {name} | {hus} |")
out.append("\n_Los cambios respecto al backlog anterior están en la hoja SUPUESTOS del Excel._\n")
(OUT / "SUPUESTOS.md").write_text("\n".join(out), encoding="utf-8")

print(len(stories), "historias,", sum(len(s["criterios"]) for s in stories), "criterios")
